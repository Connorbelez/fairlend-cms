from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/connor/Dev/fairlend-cms")
OUT = ROOT / "artifacts/fairlend-backyards-for-canadians-campaign-proposal.docx"

# Resolved preset: narrative_proposal.
# Named brand overrides: FairLend teal replaces preset heading blue; rust is used
# only for campaign kickers and emphasis. Title/cover typography follows the
# proposal_centerpiece header pattern.
TEAL = "062C2F"
TEAL_LIGHT = "E9F0EF"
RUST = "A92D17"
RUST_LIGHT = "F7EDEA"
IVORY = "F8F7F5"
INK = "171A18"
GRAY = "596260"
LIGHT_GRAY = "F4F6F9"
MID_GRAY = "D7DDDB"
WHITE = "FFFFFF"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, margins=CELL_MARGINS):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in margins.items():
        tag = "start" if edge == "start" else "end" if edge == "end" else edge
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=MID_GRAY, size=4):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def remove_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=TABLE_INDENT_DXA, borders=True):
    assert sum(widths_dxa) == CONTENT_WIDTH_DXA, widths_dxa
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_row_cant_split(row)
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if borders:
                set_cell_border(cell)
            else:
                remove_cell_borders(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    existing = tr_pr.find(qn("w:cantSplit"))
    if existing is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_keep_together(paragraph, value=True):
    paragraph.paragraph_format.keep_together = value


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left = p.add_run("FAIRLEND  /  CAMPAIGN PROPOSAL")
    set_run_font(left, size=8.5, color=TEAL, bold=True)
    right = p.add_run("                                      BACKYARDS FOR CANADIANS")
    set_run_font(right, size=8.5, color=GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    add_page_number(fp)


def add_custom_numbering(doc: Document):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(abstract_ids or [0]) + 1
    next_num = max(num_ids or [0]) + 1

    def make_abstract(abs_id: int, fmt: str, text: str, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ppr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "279")
        ppr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "290")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        lvl.append(ppr)
        if font:
            rpr = OxmlElement("w:rPr")
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), font)
            rfonts.set(qn("w:hAnsi"), font)
            rpr.append(rfonts)
            lvl.append(rpr)
        abstract.append(lvl)
        numbering.append(abstract)
        return abs_id

    bullet_abs = make_abstract(next_abs, "bullet", "•", "Arial")
    decimal_abs = make_abstract(next_abs + 1, "decimal", "%1.")
    next_steps_abs = make_abstract(next_abs + 2, "decimal", "%1.")

    def make_num(num_id: int, abs_id: int):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_id = OxmlElement("w:abstractNumId")
        abstract_id.set(qn("w:val"), str(abs_id))
        num.append(abstract_id)
        numbering.append(num)

    make_num(next_num, bullet_abs)
    make_num(next_num + 1, decimal_abs)
    make_num(next_num + 2, next_steps_abs)
    return next_num, next_num + 1, next_num + 2


def apply_num(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.208


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for style_name, size, color, before, after in (
        ("Heading 1", 16, TEAL, 18, 10),
        ("Heading 2", 13, TEAL, 12, 6),
        ("Heading 3", 12, RUST, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = rgb(TEAL)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = rgb(GRAY)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)


def add_para(doc, text="", *, align=None, before=0, after=8, size=11, color=INK,
             bold=False, italic=False, keep=False, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    if keep:
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_rich_para(doc, parts, *, after=8, before=0, align=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
    for text, kwargs in parts:
        r = p.add_run(text)
        set_run_font(r, size=kwargs.get("size", 11), color=kwargs.get("color", INK),
                     bold=kwargs.get("bold"), italic=kwargs.get("italic"))
    return p


def add_bullet(doc, text, bullet_num_id, *, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    apply_num(p, bullet_num_id)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=11, color=INK, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=11, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=INK)
    return p


def add_number(doc, text, decimal_num_id):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    apply_num(p, decimal_num_id)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_callout(doc, text, *, label=None, fill=TEAL_LIGHT, color=TEAL, size=16,
                italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, add_spacer=True):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_border(cell, color=fill, size=2)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if label:
        r = p.add_run(label.upper() + "\n")
        set_run_font(r, size=9, color=RUST, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=not italic, italic=italic)
    if add_spacer:
        after = doc.add_paragraph()
        after.paragraph_format.space_before = Pt(0)
        after.paragraph_format.space_after = Pt(4)
    return table


def add_section_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    return p


def add_label_value(doc, label, value):
    return add_rich_para(doc, [
        (label + ": ", {"bold": True, "color": TEAL}),
        (value, {}),
    ], after=5, align=WD_ALIGN_PARAGRAPH.LEFT)


def add_channel(doc, name, strategy, executions, bullet_num_id):
    add_section_heading(doc, name, 3)
    add_para(doc, strategy, after=5)
    for item in executions:
        add_bullet(doc, item, bullet_num_id)


def add_summary_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        shade_cell(cell, TEAL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=WHITE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            set_run_font(r, size=9.5, color=INK, bold=(idx == 0))
            set_cell_margins(cells[idx])
            set_cell_border(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(6)
    return table


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_major_heading(doc, text):
    p = add_section_heading(doc, text, 1)
    p.paragraph_format.page_break_before = True
    return p


def add_concept(doc, concept, bullet_num_id):
    kicker = add_para(doc, concept["kicker"].upper(), size=9.5, color=RUST, bold=True, after=3, keep=True)
    kicker.paragraph_format.page_break_before = True
    add_section_heading(doc, concept["title"], 1)
    add_callout(doc, concept["hero"], label="Hero line", fill=concept.get("fill", TEAL_LIGHT),
                color=TEAL, size=18)
    add_label_value(doc, "Strategic role", concept["role"])
    add_label_value(doc, "Primary audience", concept["audience"])
    add_para(doc, concept["idea"])

    add_section_heading(doc, "Message system", 2)
    for line in concept["messages"]:
        add_bullet(doc, line, bullet_num_id)

    add_section_heading(doc, "Per-channel strategy", 2)
    for channel in concept["channels"]:
        add_channel(doc, channel[0], channel[1], channel[2], bullet_num_id)

    add_section_heading(doc, "Physical advertising opportunities", 2)
    for item in concept["physical"]:
        add_bullet(doc, item, bullet_num_id)

    add_section_heading(doc, "Conversion path and measurement", 2)
    add_label_value(doc, "Primary CTA", concept["cta"])
    add_label_value(doc, "Landing experience", concept["landing"])
    add_label_value(doc, "Lead metrics", concept["metrics"])
    if concept.get("watch"):
        add_callout(doc, concept["watch"], label="Brand / compliance watch-out",
                    fill=RUST_LIGHT, color=RUST, size=10.5,
                    align=WD_ALIGN_PARAGRAPH.LEFT, add_spacer=False)


def build_document():
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    bullet_num_id, decimal_num_id, next_steps_num_id = add_custom_numbering(doc)

    # Cover — proposal_centerpiece pattern.
    add_para(doc, "FAIRLEND", size=11, color=RUST, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    cover_title = add_para(doc, "BACKYARDS FOR CANADIANS", size=30, color=TEAL, bold=True,
                           align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    remove_paragraph_borders(cover_title)
    add_para(doc, "A viral, values-led garden suite campaign for Toronto homeowners",
             style="Subtitle", after=10)
    add_para(doc, "Campaign Strategy & Channel Proposal", size=11, color=GRAY, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=28)
    add_callout(doc, "Your backyard could house a neighbour—not another tourist.",
                fill=TEAL, color=WHITE, size=19)

    meta = doc.add_table(rows=4, cols=2)
    rows = [
        ("Prepared for", "FairLend"),
        ("Market", "Toronto / GTA homeowner acquisition"),
        ("Campaign objective", "Create garden-suite demand through civic pride and durable homeowner economics"),
        ("Status", "Strategy proposal — creative and regulated claims require approval before publication"),
    ]
    for r_idx, (label, value) in enumerate(rows):
        cells = meta.rows[r_idx].cells
        for idx, text in enumerate((label, value)):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            set_run_font(run, size=9.5, color=TEAL if idx == 0 else INK, bold=(idx == 0))
            shade_cell(cells[idx], IVORY if idx == 0 else WHITE)
    set_table_geometry(meta, [2250, 7110])

    add_para(doc, "July 2026", size=9.5, color=GRAY,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=0)

    add_major_heading(doc, "Executive recommendation")
    add_callout(doc, "Launch one integrated campaign, not five disconnected slogans.",
                label="Recommendation", fill=TEAL_LIGHT, color=TEAL, size=16)
    add_para(doc,
             "The opportunity is to give Toronto homeowners social permission to pursue income from their property by making the decision visibly useful to the city. The campaign should not ask owners to sacrifice returns for virtue. It should make a garden suite feel like the more mature, stable and civic-minded way to monetize underused land.")
    add_para(doc,
             "The creative ideas in this proposal are strongest as a single architecture. Each line performs a different job: one earns attention, one travels socially, one names the movement, one closes the financial case, and one turns completed projects into public proof.")

    add_summary_table(doc,
                      ["Campaign layer", "Recommended line", "Job"],
                      [
                          ("Master platform", "Backyards for Canadians", "Names the movement and gives it longevity."),
                          ("Public provocation", "Homes for neighbours, not tourists.", "Creates tension, press value and instant comprehension."),
                          ("Viral hook", "Canada needs hosts. Just not the Airbnb kind.", "Subverts familiar language and invites sharing."),
                          ("Financial proof", "Build wealth. Build housing.", "Converts moral attention into homeowner self-interest."),
                          ("Local proof", "One backyard. One more Toronto home.", "Makes every completed suite measurable and visible."),
                          ("Homeowner identity", "I’m putting my backyard to work for Toronto.", "Turns participation into a public badge."),
                      ], [1875, 3450, 4035])

    add_section_heading(doc, "The central proposition", 2)
    add_callout(doc,
                "You shouldn’t have to choose between doing well and doing good. A garden suite lets you do both.",
                fill=RUST_LIGHT, color=RUST, size=15)
    add_para(doc,
             "This is the heart of the campaign. The moral proposition earns attention; the durable-income proposition earns action. FairLend makes the homeowner the protagonist: they contribute the property and required capital, while one coordinated team helps move the project from feasibility through construction financing and eligible takeout.")

    add_section_heading(doc, "What success looks like", 2)
    for item in [
        "Toronto homeowners begin to see a garden suite as a socially respected alternative to short-term rental income.",
        "The campaign produces qualified property reviews—not merely impressions or petition signatures.",
        "Every live project becomes media: lawn signs, progress stories, neighbourhood proof and a growing housing counter.",
        "FairLend owns a distinct cultural position at the intersection of housing supply, homeowner wealth and practical execution.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    add_section_heading(doc, "Strategic foundation", 1)
    add_section_heading(doc, "The cultural tension", 2)
    add_para(doc,
             "Toronto homeowners are encouraged to make productive use of their property, yet many people instinctively distrust marketing that wraps a commercial product in social-good language. This campaign earns permission by attaching the commercial offer to an observable outcome: one additional long-term home in an established neighbourhood.")
    add_para(doc,
             "The Airbnb contrast works because it is legible in one second. One model creates temporary occupancy and operational churn; the other creates durable housing and a recurring local relationship. The campaign should use that contrast as a door-opener, then quickly move to the full economic and execution case.")

    add_section_heading(doc, "The homeowner promise", 2)
    add_para(doc,
             "Turn underused land behind an existing home into a legal, income-producing rental without having to become a developer or independently coordinate the property review, design, permits, builder, construction financing, draws and takeout.")
    for item in [
        "A recurring-income asset attached to a property the homeowner already owns.",
        "A lower first-build barrier through a coordinated process and specialized financing pathway.",
        "A stable long-term tenancy proposition without guest turnover, cleaning cycles or platform dependency.",
        "A visible contribution to housing supply that the homeowner can feel proud to share.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    add_section_heading(doc, "Audience priorities", 2)
    add_summary_table(doc,
                      ["Audience", "Barrier", "Message that moves them"],
                      [
                          ("Equity-rich homeowners", "Do not know where to start", "No development experience required; begin with a property review."),
                          ("Would-be Airbnb hosts", "Assume short-term rental is the obvious income path", "Competitive, steadier income with less operational churn."),
                          ("Civic-minded owners", "Want to help but will not accept a weak investment", "Do well and do good—without pretending this is charity."),
                          ("Multi-generational families", "Need flexibility for family now and rental use later", "Create a legal, self-contained home with long-term optionality."),
                          ("Neighbourhood influencers", "Fear density will damage community character", "One carefully designed backyard home can add housing gently."),
                      ], [2200, 2800, 4360])

    add_section_heading(doc, "Viral mechanics", 2)
    for item in [
        "Conflict: neighbours versus tourists is immediately debatable and press-friendly.",
        "Identity: the homeowner is a city-builder, not merely a landlord.",
        "Participation: a pledge, property review and physical lawn sign give people something concrete to do.",
        "Visible proof: every construction site and completed suite increases the campaign counter.",
        "Utility: calculators, financing education and project roadmaps turn sharing into qualified demand.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    concepts = [
        {
            "kicker": "Concept 1 / The public provocation",
            "title": "Homes for neighbours, not tourists.",
            "hero": "Your backyard could house a neighbour—not another tourist.",
            "role": "The high-contrast launch idea. Use it to earn attention, discussion and media coverage.",
            "audience": "Toronto homeowners considering rental income, plus the broader public whose reaction creates cultural momentum.",
            "idea": "Make the housing choice unmistakable. This territory contrasts a durable local home with transient accommodation, but it never shames owners for wanting income. It redirects that ambition toward a more stable asset and a more useful outcome.",
            "messages": [
                "Homes for neighbours, not tourists.",
                "Don’t Airbnb your spare space. House someone who lives here.",
                "Toronto doesn’t need more places to visit. It needs more places to live.",
                "Competitive income without the vacancies, guest turnover, cleaning, platform dependency, or constant management of short-term rentals.",
            ],
            "channels": [
                ("Owned web and conversion", "Build a stark comparison landing page that opens with the moral contrast and immediately answers the homeowner’s financial questions.", [
                    "Above the fold: the hero line, a two-minute property-potential check and a direct CTA to book a backyard feasibility review.",
                    "Below the fold: compare long-term suite economics with short-term rental operations using transparent, property-specific assumptions rather than universal income promises.",
                    "Retarget visitors with the rational line: Build wealth. Build housing.",
                ]),
                ("Organic social", "Turn the contrast into a civic debate that people can participate in without reading a financing explainer first.", [
                    "Short video: a suitcase rolls into frame; cut to a nurse, teacher or tradesperson unlocking a garden suite. End on the hero line.",
                    "Poll: If your backyard could earn income, would you rather host 100 tourists or one Toronto household?", 
                    "Carousel: five things a long-term garden suite avoids—cleaning, reviews, guest messages, weekend turnover and platform dependency.",
                ]),
                ("Paid social and video", "Use the provocative line for thumb-stop, then qualify users around property ownership and project readiness.", [
                    "Meta homeowner audiences: 15-second comparison films and lead ads offering a free property-potential review.",
                    "YouTube pre-roll: six-second line-first spots followed by 30-second homeowner economics explainers.",
                    "Creative test: direct Airbnb language versus the softer ‘neighbour, not visitor’ variant.",
                ]),
                ("Search and intent capture", "Meet people already researching garden suites, Airbnb alternatives and rental-property income.", [
                    "Search groups: garden suite financing Toronto, laneway suite financing, Airbnb alternative, backyard rental income and secondary suite refinance.",
                    "Use neutral intent copy in search ads; reserve the provocation for the landing page and retargeting creative.",
                    "Publish a comparison guide that separates gross revenue from management effort, operating costs and vacancy assumptions.",
                ]),
                ("Earned media and community relations", "Frame the campaign as a public call to redirect homeowner entrepreneurship toward long-term housing supply.", [
                    "Launch an op-ed and media pitch around the question: What if Toronto’s next housing program started in its backyards?", 
                    "Offer real project walkthroughs and qualified experts rather than generic housing-crisis commentary.",
                    "Invite neighbourhood associations to a ‘neighbours, not tourists’ backyard housing forum.",
                ]),
                ("Email, CRM and partners", "Use the cultural hook to reopen dormant homeowner leads, then move quickly into feasibility and financing.", [
                    "Email subject: Your backyard could house a neighbour—and pay you every month.",
                    "Broker/realtor kit: one-page comparison, referral QR code and compliant explanation of eligible financing pathways.",
                    "Builder follow-up: give every site-visit lead a co-branded project roadmap.",
                ]),
            ],
            "physical": [
                "TTC shelters and bus tails in homeowner-dense neighbourhoods: ‘Homes for neighbours, not tourists.’ QR: See what your backyard could become.",
                "Split-panel street posters: a rolling suitcase on one side, a front-door key on the other; minimal copy for instant comprehension.",
                "Direct mail: an oversized door-hanger-shaped card reading ‘Your backyard has another front door in it.’",
                "Community newspaper and real-estate-window ads: longer-form comparison creative with a property-review offer.",
                "Construction hoarding: ‘This backyard is becoming a home for a neighbour.’ Add the live campaign counter and QR code.",
            ],
            "cta": "See what your backyard could become.",
            "landing": "Provocation → long-term-versus-short-term comparison → property checker → consultation booking.",
            "metrics": "Share rate, earned mentions, branded search lift, property-check starts, qualified review bookings and sentiment split.",
            "watch": "Do not imply every garden suite will outperform an Airbnb. Any income comparison must use disclosed, property-specific assumptions and distinguish gross revenue from net operating effort and costs.",
        },
        {
            "kicker": "Concept 2 / The viral social hook",
            "title": "Canada needs hosts. Just not the Airbnb kind.",
            "hero": "Canada needs hosts. Just not the Airbnb kind.",
            "role": "The most inherently shareable line. Use it for social-first creative, earned media and cultural conversation.",
            "audience": "Homeowners, younger family members who influence them, housing advocates, urbanists and local media.",
            "idea": "Reclaim the word ‘host.’ The homeowner still hosts, but the relationship is long-term and rooted in community. The linguistic twist gives the campaign humour and edge while leaving room for an optimistic story about new families, workers and neighbours.",
            "messages": [
                "Canada needs hosts. Just not the Airbnb kind.",
                "Host a household. Grow a neighbourhood.",
                "Be the reason someone can live closer to work.",
                "Build wealth. House a neighbour.",
            ],
            "channels": [
                ("Owned web and conversion", "Create a social-native campaign page with a pledge, share card and immediate property-review path.", [
                    "Let visitors generate a share card: ‘I’d host a neighbour in my backyard.’",
                    "Follow the share action with the practical steps, eligibility caveats and a consultation CTA.",
                    "Feature short homeowner profiles answering why they chose a long-term tenant over guest turnover.",
                ]),
                ("Organic social", "Build repeatable formats around the redefinition of ‘host’.", [
                    "Street interview series: ‘What kind of host does Toronto need?’",
                    "Creator prompt: show the person who keeps your neighbourhood running, then ask what a shorter commute would mean to them.",
                    "Comment-friendly static posts that leave the line largely unexplained and allow the audience to debate it.",
                ]),
                ("Paid social and video", "Use the punchline as a pattern interrupt and the homeowner outcome as the qualification layer.", [
                    "Six-second bumper: ‘Canada needs hosts.’ Beat. ‘Just not the Airbnb kind.’ FairLend + CTA.",
                    "Creator whitelisting with local real-estate, construction and personal-finance voices.",
                    "Sequential retargeting: joke → homeowner story → property-potential offer.",
                ]),
                ("Search and intent capture", "Do not force the joke into every high-intent ad. Use it as a recognizable brand asset once intent has been captured.", [
                    "Bid on host and Airbnb-adjacent queries only where policy and relevance permit.",
                    "Build an SEO article around long-term rental alternatives for Toronto homeowners.",
                    "Retarget comparison-guide readers with the viral line and a feasibility review.",
                ]),
                ("Earned media and community relations", "Stage a public redefinition of hosting around housing contribution.", [
                    "Recruit a small founding group of ‘Toronto’s new hosts’ and announce the first pledged backyards.",
                    "Pitch morning radio and local television on the provocative language and the real homeowners behind it.",
                    "Partner with employers, hospitals or trade groups on stories about workers who need homes near work—without turning tenants into props.",
                ]),
                ("Email, CRM and partners", "Give advocates and referral partners a line they will actually forward.", [
                    "Email subject: Canada needs a different kind of host.",
                    "Referral card: Know a homeowner with room for one more Toronto home?", 
                    "Partner toolkit with approved copy variants ranging from provocative to conservative.",
                ]),
            ],
            "physical": [
                "Wild-posting series: ‘HOST WANTED’ in oversized type, revealed below as ‘Backyard required. Tourism experience not necessary.’",
                "Coffee sleeves and takeout bags near homeowner neighbourhoods: ‘A different kind of host keeps Toronto running.’",
                "GO/TTC station dominations near commuter corridors: ‘Host someone closer to work.’",
                "Neighbourhood event placards and photo walls: ‘Toronto’s new hosts.’",
                "Realtor open-house inserts: ‘The next income suite may be behind this house.’",
            ],
            "cta": "Become a different kind of host.",
            "landing": "Pledge/share moment → homeowner stories → how the project works → property review.",
            "metrics": "Organic shares, creator remix volume, direct traffic, pledge completions, assisted conversions and partner referrals.",
            "watch": "Keep the word ‘Canada’ inclusive. Stories and casting should visibly include newcomers and different household types. The campaign must never suggest citizenship-based tenant preference or exclusion.",
        },
        {
            "kicker": "Concept 3 / The movement platform",
            "title": "Backyards for Canadians",
            "hero": "I’m putting my backyard to work for Toronto.",
            "role": "The durable campaign platform. Use it to organize pledges, partnerships, events, project stories and the public impact counter.",
            "audience": "Civic-minded homeowners, neighbourhood leaders, employers, builders, realtors, brokers and public-interest partners.",
            "idea": "Turn individual projects into a visible movement. Homeowners should be able to signal participation before, during and after construction. The platform becomes more credible with every signed lawn, funded project and completed home.",
            "messages": [
                "Backyards for Canadians.",
                "I’m putting my backyard to work for Toronto.",
                "Your backyard can help solve the housing crisis.",
                "A backyard shouldn’t sit empty during a housing crisis.",
                "Build wealth. House a neighbour.",
            ],
            "channels": [
                ("Owned web and conversion", "Build the campaign hub as a live movement dashboard rather than a conventional product page.", [
                    "Map pledged, reviewed, financed, under-construction and completed backyards at an appropriate privacy-preserving geographic level.",
                    "Show a live total: backyards pledged, projects in review and long-term homes created.",
                    "Give each participant a unique share page and QR code tied to their project stage.",
                ]),
                ("Organic social", "Make participation visible and repeatable.", [
                    "Weekly ‘Backyard of the Week’ stories focused on the homeowner’s motivation and project learning.",
                    "Milestone posts: feasibility confirmed, permit submitted, foundation poured, tenant welcomed.",
                    "UGC template: ‘I’m putting my backyard to work for ___.’ Participants fill in Toronto, my family, a neighbour or the next generation.",
                ]),
                ("Paid social and video", "Use social proof and geographic proximity to reduce the perceived novelty of building.", [
                    "Neighbourhood creative: ‘Three backyards in East York are already in review.’ Use only verified counts.",
                    "Test founder-style videos against polished campaign films.",
                    "Retarget pledged homeowners with process education and booking prompts so the movement produces projects.",
                ]),
                ("Search and intent capture", "Own the campaign name and connect broad housing interest to a concrete homeowner action.", [
                    "Create campaign-name and neighbourhood landing pages with a consistent property-check CTA.",
                    "Publish practical guides for feasibility, permitting, construction draws, takeout and landlord readiness.",
                    "Use structured FAQs to answer whether a garden or laneway suite may fit a property before asking for a call.",
                ]),
                ("Earned media and community relations", "Give media a measurable movement with recurring milestones.", [
                    "Launch with a founding pledge and a transparent first-year housing target only after capacity and assumptions are validated.",
                    "Create an annual ‘State of Toronto’s Backyards’ report based on campaign pipeline data.",
                    "Host open-backyard tours with builders, planners and financing specialists.",
                ]),
                ("Email, CRM and partners", "Turn the platform into a shared acquisition layer across the housing ecosystem.", [
                    "Partner-specific pledge pages and QR codes for builders, realtors, mortgage professionals and community groups.",
                    "Milestone email sequence that tells participants what happens next and what documentation to prepare.",
                    "Neighbour referral loop: when one property enters review, invite nearby owners to a group information session.",
                ]),
            ],
            "physical": [
                "Signature lawn sign during construction: ‘Another Toronto home is being built here.’ Include the campaign mark, project counter and QR code.",
                "Permanent completion plaque: ‘One backyard. One more Toronto home.’ Offer only with homeowner consent.",
                "Neighbourhood counter boards at community events, builder showrooms and partner offices.",
                "Branded project fencing and trade-site signage that evolves with each construction milestone.",
                "A mobile ‘Backyard Housing Lab’ pop-up at street festivals with parcel sketches, financing education and on-site property-review booking.",
            ],
            "cta": "Put your backyard to work for Toronto.",
            "landing": "Impact counter → participant stories → pledge → property review → project-stage updates.",
            "metrics": "Pledges, pledge-to-review rate, partner-sourced reviews, project-stage progression, verified homes created and neighbourhood referral rate.",
            "watch": "A pledge counter without completed projects will become empty virtue signalling. Publish stage definitions, verified counts and dates; never imply that a pledge equals a financed or completed home.",
        },
        {
            "kicker": "Concept 4 / The rational conversion platform",
            "title": "Build wealth. Build housing.",
            "hero": "Build wealth. Build housing.",
            "role": "The financially grounded conversion idea. Use it wherever homeowner economics, financing and process clarity matter most.",
            "audience": "Equity-rich homeowners, financially motivated households, mortgage clients, real-estate investors and referral professionals.",
            "idea": "Make the dual benefit explicit. This concept is less provocative and more scalable across regulated advertising, search, CRM and partner channels. It reassures owners that the social outcome is not charity—it is created by a well-structured property investment.",
            "messages": [
                "Build wealth. Build housing.",
                "Your property can do more than appreciate.",
                "Turn unused land into someone’s first home.",
                "Create reliable, long-term income—and one more place to live.",
                "You don’t need to be a developer. You need the right team and a viable property.",
            ],
            "channels": [
                ("Owned web and conversion", "Lead with a transparent project equation and an easy first step.", [
                    "Interactive property-potential tool: lot basics, project range, likely next questions and consultation booking.",
                    "Explain the coordinated workstreams: feasibility, design/permits, builder, construction financing, draws, completion and eligible takeout.",
                    "Use scenario ranges and disclosed assumptions rather than single-point income promises.",
                ]),
                ("Organic social", "Teach the economics in plain language while keeping the social outcome visible.", [
                    "Carousel: seven costs people forget when comparing Airbnb with a long-term suite.",
                    "Whiteboard videos: how construction draws and takeout can fit together.",
                    "Before/after parcel sketches showing how underused land becomes a self-contained rental home.",
                ]),
                ("Paid social and video", "Use homeowner self-interest to generate qualified leads, with the civic line as differentiation.", [
                    "Lead ads: ‘Could your backyard support a long-term rental home?’",
                    "Finance explainer videos targeted to homeowners, mortgage intenders and renovation audiences.",
                    "Retarget property-tool abandoners with a consultation offer and process reassurance.",
                ]),
                ("Search and intent capture", "This should be the dominant search territory because it maps directly to active homeowner questions.", [
                    "Build tightly matched pages for garden suite financing, laneway suite financing, secondary suite refinance and construction draws.",
                    "Use compliant qualifiers around amortization, loan-to-value, approval and program eligibility.",
                    "Offer a downloadable project-planning checklist as the mid-intent conversion.",
                ]),
                ("Earned media and community relations", "Contribute useful homeowner economics to housing coverage instead of relying only on moral messaging.", [
                    "Publish anonymized project scenarios showing the variables that determine viability.",
                    "Provide experts for stories about financing gentle density and reducing first-build friction.",
                    "Run homeowner clinics with builders and mortgage professionals in target neighbourhoods.",
                ]),
                ("Email, CRM and partners", "Give leads a disciplined education sequence that reduces fear and surfaces readiness.", [
                    "Sequence: property fit → project budget → financing path → permits/build → long-term rental readiness.",
                    "Partner deck for accountants, planners, realtors and brokers explaining who is and is not a good fit.",
                    "Re-engagement email: Your property may be worth more as a project than as unused land.",
                ]),
            ],
            "physical": [
                "Targeted direct mail with a parcel-sketch visual: ‘There may be another income-producing home behind yours.’",
                "Home-improvement retail placements near contractor desks and design centres: ‘You renovate rooms. What if you built an address?’",
                "Mortgage and realtor office displays with a take-one project checklist and trackable QR code.",
                "Neighbourhood seminar posters: ‘Garden suite feasibility + financing clinic.’",
                "Construction vehicle and trailer graphics: ‘Build wealth. Build housing.’ paired with a completed-suite image.",
            ],
            "cta": "Check your property’s potential.",
            "landing": "Property inputs → viability education → transparent project equation → consultation booking.",
            "metrics": "Cost per qualified property review, tool completion, consultation show rate, viable-property rate, application rate and project conversion.",
            "watch": "‘Financing available over 30 years’ is too broad as public copy. Safer working language: ‘Eligible financing may amortize up to 30 years, subject to approval and program requirements.’ Longer amortization can reduce monthly payments but increases total interest paid.",
        },
        {
            "kicker": "Concept 5 / The proof and neighbourhood platform",
            "title": "One backyard. One more Toronto home.",
            "hero": "One backyard. One more Toronto home.",
            "role": "The proof system. Use it to make gentle density tangible, local and cumulative.",
            "audience": "Homeowners who need social proof, neighbours worried about change, municipal and community stakeholders, and local press.",
            "idea": "The housing crisis is abstract; one new home is concrete. This territory turns every project into a unit of progress and makes the campaign naturally hyperlocal. It is the strongest system for site signage, neighbourhood storytelling and impact reporting.",
            "messages": [
                "One backyard. One more Toronto home.",
                "One lot can hold more possibility.",
                "Another Toronto home is being built here.",
                "Turn unused land into someone’s first home.",
                "Small build. Real home. Lasting impact.",
            ],
            "channels": [
                ("Owned web and conversion", "Organize proof geographically and by project stage.", [
                    "Neighbourhood pages with verified counts, anonymized project examples and local considerations.",
                    "A visual project timeline from backyard review to tenant-ready home.",
                    "A counter that distinguishes reviews, financed builds, construction starts and completed homes.",
                ]),
                ("Organic social", "Use progress content instead of repeatedly explaining the campaign idea.", [
                    "Monthly ‘one more home’ reveal videos.",
                    "Construction time-lapse with one useful lesson at every milestone.",
                    "Neighbourhood series: what one garden suite can mean for a teacher, caregiver, new family or aging parent.",
                ]),
                ("Paid social and video", "Use nearby proof to reduce perceived risk and increase consultation quality.", [
                    "Geo-tailored creative: ‘See how one East York backyard became one more home.’",
                    "Case-study ads with verified timelines, project scope and homeowner motivation.",
                    "Sequential video: empty yard → construction → finished front door → long-term home.",
                ]),
                ("Search and intent capture", "Connect neighbourhood proof with high-intent feasibility queries.", [
                    "Local pages for Toronto neighbourhoods where planning guidance and project examples are genuinely distinct.",
                    "Case-study schema and internal links from financing guides to completed-project stories.",
                    "Search ad CTA: See a Toronto garden suite project from first review to takeout.",
                ]),
                ("Earned media and community relations", "Create recurring, evidence-led local news moments.", [
                    "Invite local reporters to verified project milestones rather than issuing generic campaign announcements.",
                    "Release a quarterly counter with methodology and stage definitions.",
                    "Host small open-house tours for neighbours before broad public events.",
                ]),
                ("Email, CRM and partners", "Use completed proof to move cautious homeowners through the funnel.", [
                    "Case-study emails segmented by lot type, homeowner motivation and project stage.",
                    "Builder and realtor leave-behind: one project, one page, one QR code.",
                    "Neighbour sequence triggered by a nearby project information session, subject to consent and privacy controls.",
                ]),
            ],
            "physical": [
                "Numbered lawn signs: ‘Toronto home #___ is being built here’ only after the counting methodology is established.",
                "Construction-fence progress panels that change at feasibility, permit, build and completion stages.",
                "Hyperlocal mail around consenting project sites: ‘One more home is coming to the neighbourhood. Here’s how it works.’",
                "TTC platform sequence: panel one ‘One backyard.’ Panel two ‘One more Toronto home.’ Panel three property-review CTA.",
                "Permanent or semi-permanent completion markers that homeowners can opt into as a civic badge.",
            ],
            "cta": "Create one more Toronto home.",
            "landing": "Local proof → verified case study → project process → property review.",
            "metrics": "Case-study engagement, neighbourhood traffic, local referral rate, review-to-viability rate and verified completed homes.",
            "watch": "Count only legal, self-contained long-term homes that reach the published project stage. Do not use tenants’ identities, occupations or personal stories without informed consent and fair compensation where appropriate.",
        },
    ]

    for concept in concepts:
        add_concept(doc, concept, bullet_num_id)

    add_major_heading(doc, "Cross-channel operating model")
    add_para(doc,
             "Each concept can run independently, but the recommended system deliberately hands the prospect from cultural attention to financial proof. Every channel should have one primary job and one measurable next action.")
    add_summary_table(doc,
                      ["Channel", "Primary job", "Recommended campaign layer", "Conversion"],
                      [
                          ("PR / earned", "Create debate and legitimacy", "Homes for neighbours / Canada needs hosts", "Campaign hub visits"),
                          ("Organic social", "Earn shares and participation", "Canada needs hosts / pledge", "Pledge or property check"),
                          ("Paid social / video", "Qualify and retarget", "Provocation → wealth/housing proof", "Property review"),
                          ("Search", "Capture active demand", "Build wealth. Build housing.", "Consultation booking"),
                          ("Email / CRM", "Educate and progress", "One team / first review to takeout", "Documents + appointment"),
                          ("Partners", "Borrow trust and referrals", "Backyards for Canadians", "Trackable referral"),
                          ("Physical / OOH", "Make the movement locally visible", "One backyard. One more home.", "QR scan / event"),
                      ], [1500, 2250, 3370, 2240])

    add_section_heading(doc, "The conversion spine", 2)
    steps = [
        "Attention: encounter the provocation in social, press or the street.",
        "Identification: see a homeowner, property or neighbourhood that feels familiar.",
        "Utility: complete a property-potential check or download the planning checklist.",
        "Qualification: book a review and provide the inputs needed to assess fit.",
        "Progression: receive a clear roadmap for feasibility, capital, permits, construction and eligible takeout.",
        "Advocacy: share the pledge, lawn sign or verified project milestone.",
    ]
    for item in steps:
        add_number(doc, item, decimal_num_id)

    add_section_heading(doc, "Measurement infrastructure", 2)
    for item in [
        "Unique QR codes by physical placement, partner, neighbourhood and creative concept.",
        "Consistent UTMs and CRM source fields through property review, consultation, viability and application stages.",
        "Separate campaign counters for pledges, reviews, financed projects, starts and completed homes.",
        "Creative-level sentiment tracking so provocation is evaluated against qualified demand—not engagement alone.",
        "Call tracking and self-reported attribution for community events, direct mail and partner referrals.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    add_major_heading(doc, "Physical media master plan")
    add_para(doc,
             "Physical media is not an accessory to this campaign; it is the proof layer. A backyard-housing movement becomes credible when people can see projects in their own neighbourhoods. Prioritize placements that either sit on the property itself or reach homeowners in a geographically accountable way.")
    add_summary_table(doc,
                      ["Medium", "Best concept", "Execution", "Why it earns its keep"],
                      [
                          ("Lawn signs", "One more Toronto home", "Milestone-based signs on consenting project sites", "Every project becomes a local endorsement and lead source."),
                          ("Construction hoarding", "Backyards for Canadians", "Progress story + live counter + QR", "Long dwell time and immediate proof."),
                          ("TTC / shelter", "Neighbours, not tourists", "Minimal copy, visual contrast, short URL", "Creates cultural reach and press photography."),
                          ("Direct mail", "Build wealth. Build housing.", "Parcel-sketch mailer to selected postal walks", "Reaches property owners with enough room to explain the offer."),
                          ("Door hangers", "One backyard", "Neighbour information + event invite", "Useful around active sites when handled transparently."),
                          ("Retail / showroom", "Build wealth", "Checklist racks at builders, design centres and realtor offices", "Catches homeowners during renovation and property decisions."),
                          ("Community events", "Backyards for Canadians", "Mobile housing lab + on-site booking", "Turns abstract interest into practical review."),
                          ("Completion plaques", "One more Toronto home", "Optional permanent civic marker", "Creates enduring recognition and neighbour curiosity."),
                      ], [1600, 2100, 3160, 2500])

    add_section_heading(doc, "Signature physical asset: the lawn sign system", 2)
    for item in [
        "Stage 1 — Review: ‘This backyard may become one more Toronto home.’",
        "Stage 2 — Approved / proceeding: ‘Another Toronto home is being planned here.’",
        "Stage 3 — Construction: ‘Another Toronto home is being built here.’",
        "Stage 4 — Completion: ‘One backyard. One more Toronto home.’",
        "Every sign uses a stage-specific QR code and never overstates project certainty.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    add_section_heading(doc, "Placement principles", 2)
    for item in [
        "Choose homeowner-dense postal walks and neighbourhoods based on property fit, project capacity and serviceability—not broad citywide reach alone.",
        "Keep provocative OOH copy extremely short; use the QR destination to add nuance.",
        "Treat active sites as community communications: identify what is happening, expected stages and where neighbours can learn more.",
        "Obtain owner and site permissions, comply with municipal sign rules and avoid any suggestion of government endorsement.",
        "Use distinct tracking for every medium so physical spend is evaluated on qualified reviews, not scans alone.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    add_major_heading(doc, "90-day launch strategy")
    add_summary_table(doc,
                      ["Phase", "Weeks", "What ships", "Decision gate"],
                      [
                          ("1. Prove", "1–2", "Claim substantiation, campaign identity, landing prototype, property checker, CRM/QR instrumentation, founding homeowner recruitment", "Can FairLend responsibly support the offer and measure qualified demand?"),
                          ("2. Pilot", "3–6", "Three social concepts, search capture, targeted mail, first lawn signs, one community clinic", "Which message produces viable properties—not only engagement?"),
                          ("3. Activate", "7–10", "PR launch, creator content, transit test, partner toolkit, public counter and project stories", "Does the integrated architecture improve branded search and review volume?"),
                          ("4. Scale", "11–13", "Shift spend to winning neighbourhoods and creative; expand site signage and referrals", "Can operations absorb the lead and project pipeline without degrading experience?"),
                      ], [1450, 850, 3970, 3090])

    add_section_heading(doc, "Minimum viable launch package", 2)
    for item in [
        "One campaign hub with property-potential intake, transparent stage definitions and compliant financing language.",
        "Three launch films: provocation, homeowner story and rational process/economics.",
        "Six static/social units across the campaign architecture.",
        "A four-stage lawn sign system and one targeted direct-mail piece.",
        "Partner referral kit for builders, realtors and mortgage professionals.",
        "Founding homeowner cohort with consented stories and a clear operational path.",
        "Measurement dashboard from impression and QR scan through viable property and completed home.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    add_section_heading(doc, "Budget allocation logic", 2)
    add_para(doc,
             "Allocate the pilot by learning objective rather than a premature fixed media split. Protect budget for production and attribution. Search captures existing demand; social and physical media create demand and cultural permission. Scale only after testing qualified-property rate and operational capacity.")

    add_major_heading(doc, "Claims, compliance and brand safety")
    add_callout(doc,
                "The campaign can be morally bold without being financially sloppy.",
                label="Operating principle", fill=RUST_LIGHT, color=RUST, size=15)
    add_section_heading(doc, "Income comparison", 2)
    add_para(doc,
             "Do not publicly promise that every garden suite will make ‘just as good an income as Airbnb.’ The defensible proposition is stable, potentially competitive long-term income with less turnover and platform dependency. Any numeric comparison must disclose property-specific assumptions, rate, amortization, mortgage amount, insurance premium, taxes, utilities, maintenance, vacancy, contingency, management and short-term-rental operating costs.")

    add_section_heading(doc, "Financing language", 2)
    for item in [
        "Preserve the working concept ‘No development experience required,’ but do not imply that FairLend eliminates owner responsibilities or guarantees project execution.",
        "Replace broad public statements such as ‘Financing available over 30 years’ with qualified language tied to the applicable eligible program.",
        "Do not combine maximum loan-to-value and maximum amortization figures from different eligibility paths as though they are one guaranteed structure.",
        "All financing is subject to property fit, borrower qualification, underwriting, approval, program rules and available capital.",
        "Any ad containing an interest rate, payment amount or non-interest charge requires the applicable APR and term disclosures with appropriate prominence.",
    ]:
        add_bullet(doc, item, bullet_num_id)

    add_section_heading(doc, "Housing and inclusion", 2)
    for item in [
        "‘Canadians’ must include the full lived reality of Canada: newcomers, long-time residents, different family structures and people across income levels.",
        "Do not suggest citizenship, occupation or family-status preferences in tenant selection.",
        "Avoid using renters as moral props. Obtain informed consent for stories and make the homeowner/project—not a vulnerable tenant—the default subject.",
        "Use ‘affordable’ only when the campaign defines and substantiates the standard. Otherwise prefer ‘high-quality long-term housing’ or ‘a home closer to work and community.’",
    ]:
        add_bullet(doc, item, bullet_num_id)

    add_section_heading(doc, "Required approval gate", 2)
    add_para(doc,
             "Before publication, every regulated asset should be reviewed by FairLend’s principal broker/compliance function. The final creative should clearly display the required authorized name and licensing information, avoid government-approval implications, and maintain an asset register showing approved copy, disclosures, dates and channels.")

    add_major_heading(doc, "Complete original copy library")
    add_para(doc,
             "This appendix preserves every campaign line and copy concept from the original strategy response. The copy is intentionally retained before channel adaptation; regulated and performance claims remain subject to substantiation and approval.")

    add_section_heading(doc, "Core framing", 2)
    add_callout(doc, "Your backyard could house a neighbour—not another tourist.", fill=TEAL, color=WHITE, size=17)
    add_label_value(doc, "Campaign platform", "Backyards for Canadians")

    add_section_heading(doc, "Core message", 2)
    add_para(doc, "Toronto doesn’t need more places to visit. It needs more places to live.", bold=True)
    add_para(doc,
             "Build a garden suite. Give a working Canadian or young family a high-quality home near their job, school, and community—and turn your property into reliable, long-term income.")
    add_para(doc, "No development experience required. Financing available over 30 years.")

    add_section_heading(doc, "Headline bank", 2)
    original_headlines = [
        "Homes for neighbours, not tourists.",
        "Your backyard can help solve the housing crisis.",
        "Don’t Airbnb your spare space. House someone who lives here.",
        "A backyard shouldn’t sit empty during a housing crisis.",
        "Build wealth. Build housing.",
        "One backyard. One more Toronto home.",
        "Your property can do more than appreciate.",
        "Turn unused land into someone’s first home.",
        "Canada needs hosts. Just not the Airbnb kind.",
    ]
    for item in original_headlines:
        add_bullet(doc, item, bullet_num_id)
    add_para(doc, "The strongest viral line is probably:", italic=True, color=GRAY)
    add_callout(doc, "Canada needs hosts. Just not the Airbnb kind.", fill=RUST_LIGHT, color=RUST, size=16)
    add_para(doc,
             "It subverts familiar language, invites discussion, and positions FairLend as offering a socially useful alternative rather than merely selling debt.")

    add_section_heading(doc, "Short manifesto", 2)
    manifesto = [
        "Toronto homeowners are sitting on something the city desperately needs: space.",
        "A garden suite can become a high-quality home for a nurse, teacher, tradesperson, new Canadian, or young family—close to the work and community that depend on them.",
        "It can also provide homeowners with dependable, long-term rental income and increase the productive value of their property.",
        "You don’t need to be a developer. You don’t need construction experience. With long-term financing and the right team, an unused backyard can become part of the solution.",
    ]
    for para in manifesto:
        add_para(doc, para)
    add_callout(doc, "Build wealth. House a neighbour.", fill=TEAL_LIGHT, color=TEAL, size=16)

    add_section_heading(doc, "The social mechanism", 2)
    add_para(doc, "The campaign becomes shareable when it gives homeowners a visible commitment:")
    add_callout(doc, "I’m putting my backyard to work for Toronto.", fill=RUST_LIGHT, color=RUST, size=16)
    for item in [
        "A physical lawn sign during construction: ‘Another Toronto home is being built here.’",
        "A digital ‘Backyard Housing Pledge’ badge.",
        "A shareable estimate: ‘This property could create one home for the next 30+ years.’",
        "A neighbourhood counter showing how many homes participating owners have created.",
        "Profiles of the workers and families garden suites make room for—without exploiting individual tenants.",
    ]:
        add_bullet(doc, item, bullet_num_id)
    add_para(doc,
             "The lawn signs are particularly strong. Every project becomes a local advertisement and a public declaration of civic participation.")

    add_section_heading(doc, "Critical positioning adjustment", 2)
    add_para(doc,
             "I would not publicly promise homeowners they can make ‘just as good an income as Airbnb’ unless FairLend can substantiate it with Toronto-specific numbers.")
    add_para(doc, "The more defensible—and probably more persuasive—claim is:")
    add_callout(doc,
                "Competitive income without the vacancies, guest turnover, cleaning, platform dependency, or constant management of short-term rentals.",
                fill=TEAL_LIGHT, color=TEAL, size=14)
    add_para(doc,
             "That shifts the comparison from gross nightly revenue to stable, lower-effort, long-term economics. It also protects the campaign’s credibility: the moral proposition gets attention, but the financial case closes the homeowner.")
    add_para(doc, "The campaign should never sound like charity. The underlying proposition is stronger:")
    add_callout(doc,
                "You shouldn’t have to choose between doing well and doing good. A garden suite lets you do both.",
                fill=RUST_LIGHT, color=RUST, size=14)
    add_para(doc,
             "That is the real idea: turn participation in Toronto’s housing solution into something homeowners can feel proud of—and profit from.")

    add_major_heading(doc, "Decision and immediate next actions")
    add_callout(doc,
                "Approve the integrated architecture and pilot it against qualified-property outcomes.",
                label="Decision requested", fill=TEAL, color=WHITE, size=15)
    next_steps = [
        "Select the master public name: Backyards for Canadians, with Backyards for Toronto as the local activation variant.",
        "Approve ‘Homes for neighbours, not tourists’ and ‘Canada needs hosts’ for controlled creative testing.",
        "Commission the claim-substantiation and compliance pass before producing financial comparison assets.",
        "Recruit three to five founding homeowners or active projects that can supply authentic proof.",
        "Build the campaign hub, property checker, CRM attribution and stage-based public counter.",
        "Produce the launch creative and physical lawn-sign system as one coherent identity.",
        "Run the 90-day pilot, then scale based on viable-property rate, project capacity and brand sentiment.",
    ]
    for item in next_steps:
        add_number(doc, item, next_steps_num_id)

    add_para(doc, "Build wealth. Build housing. House a neighbour.", size=14, color=TEAL,
             bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=18, after=0)

    # Document properties.
    props = doc.core_properties
    props.title = "Backyards for Canadians — Campaign Strategy & Channel Proposal"
    props.subject = "Viral, values-led garden suite campaign for Toronto homeowners"
    props.author = "FairLend"
    props.keywords = "FairLend, garden suites, Toronto, campaign strategy, physical advertising"
    props.comments = "Strategy proposal; creative and regulated claims require approval before publication."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
